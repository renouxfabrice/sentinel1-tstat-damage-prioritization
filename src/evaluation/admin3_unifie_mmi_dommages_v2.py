# -*- coding: utf-8 -*-
"""Même calcul unifié que la version précédente, avec la séparation des zones couvertes et non couvertes par les nuages et un comptage par produit étendu.

Les chemins d'entree et de sortie sont declares dans le bloc CONFIGURATION
ci-dessous, relativement aux racines de donnees fixees dans `config.py`.
"""

import os
import sys
import csv
import numpy as np
from config import DONNEES

try:
    from osgeo import ogr, osr, gdal
except ImportError:
    sys.exit("osgeo (gdal/ogr) requis.")

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl requis : pip install openpyxl --break-system-packages")

ogr.UseExceptions()
gdal.UseExceptions()

# =============================================================================
# CONFIGURATION — UNE SEULE FOIS, PARTAGEE PAR LES DEUX ANALYSES
# =============================================================================

AOI_PATH = os.path.join(DONNEES, r"VENEZUELA\EMS\EMSR884_products\EMSR884_ALL\EMSR884_areOfinterestA_v3.gpkg")
AOI_AREA_FIELD = "area"

ADMIN3_PATH = os.path.join(DONNEES, r"VENEZUELA\HDX\ven_admin_boundaries.geojson\ven_admin3.geojson")
ADMIN3_ID_FIELD = 'adm3_pcode'

BUILDINGS_OVERTURE_PATH = os.path.join(DONNEES, r"VENEZUELA\Building\hot_eq_overture_ven_buildings_overture_gpkg\buildings.gpkg")
BUILDINGS_OVERTURE_FIELD_OVERRIDE = None

BUILDINGS_OSM_PATH = os.path.join(DONNEES, r"VENEZUELA\Building\hot_eq_osm_ven_buildings_osm_gpkg\buildings.gpkg")
BUILDINGS_OSM_FIELD_OVERRIDE = [
    'id', 'name', 'name_en', 'name_es', 'building', 'building_levels',
    'building_materials', 'addr_full', 'addr_housenumber', 'addr_street',
    'addr_city', 'office', 'source', 'adm0_pcode', 'adm0_name',
    'adm1_pcode', 'adm1_name', 'adm2_pcode', 'adm2_name', 'adm3_pcode',
    'adm3_name', 'adm4_pcode', 'adm4_name', 'name_latin',
]
POPULATION_PATH = os.path.join(DONNEES, r"VENEZUELA\HDX\ven_pop_2026_CN_100m_R2025A_v1.tif")
GHS_BUILT_PATH = os.path.join(DONNEES, r"VENEZUELA\HDX\GHS_BUILT_S_E2025_GLOBE_R2023A_54009_100_V1_0_R8_C12\GHS_BUILT_S_E2025_GLOBE_R2023A_54009_100_V1_0_R8_C12.tif")

MMI_SHP_PATH = os.path.join(DONNEES, r"VENEZUELA\UNGSC\M 7.5 - 17 km W of Catia La Mar_Venezuela\mi.shp")
MMI_VALUE_FIELD = "PARAMVALUE"
MMI_BINS = [
    (float('-inf'), 1.5, 'I'),
    (1.5, 3.5, 'II-III'),
    (3.5, 4.5, 'IV'),
    (4.5, 5.5, 'V'),
    (5.5, 6.5, 'VI'),
    (6.5, 7.5, 'VII'),
    (7.5, 8.5, 'VIII'),
    (8.5, 9.5, 'IX'),
    (9.5, float('inf'), 'X+'),
]

NOT_ANALYSED_PATHS = [
    os.path.join(DONNEES, r"VENEZUELA\EMS\EMSR884_products\EMSR884_AOI02_GRA_MONIT01_v2\EMSR884_AOI02_GRA_MONIT01_notAnalysedA_v1.shp"),
    os.path.join(DONNEES, r"VENEZUELA\EMS\EMSR884_products\EMSR884_AOI06_GRA_MONIT01_v2\EMSR884_AOI06_GRA_MONIT01_notAnalysedA_v1.shp"),
    os.path.join(DONNEES, r"VENEZUELA\EMS\EMSR884_products\EMSR884_AOI08_GRA_MONIT01_v2\EMSR884_AOI08_GRA_MONIT01_notAnalysedA_v1.shp"),
    os.path.join(DONNEES, r"VENEZUELA\EMS\EMSR884_products\EMSR884_AOI12_GRA_MONIT01_v3\EMSR884_AOI12_GRA_MONIT01_notAnalysedA_v1.shp"),
]

DAMAGE_CONFIG = {
    'chatmap': {
        'kind': 'points',
        'path': os.path.join(DONNEES, r"VENEZUELA\HOTOSM\chatmap\damage_validation_buildings_venezuela_points.geojson"),
        'class_field': 'damaged', 'damaged_values': ['significant', 'minimal'],
        'destroyed_values': ['complete'],
    },
    'disha': {
        'kind': 'points_csv',
        'path': os.path.join(DONNEES, r"VENEZUELA\DISHA\Caracas\nwcaracas_final_inference_result.csv"),
        'class_field': 'damaged', 'damaged_values': ['TRUE', 'True', True], 'destroyed_values': [],
    },
    'ems': {
        'kind': 'points',
        'path': os.path.join(DONNEES, r"VENEZUELA\EMS\EMSR884_products\EMSR884_ALL\EMSR884_builtUpP.gpkg"),
        'class_field': 'damage_gra', 'damaged_values': ['Possibly damaged', 'Damaged'],
        'destroyed_values': ['Destroyed'],
    },
    'fair': {
        'kind': 'points',
        'path': os.path.join(DONNEES, r"VENEZUELA\Fair\ALL\damage_polygones_fair.gpkg"),
        'class_field': 'damage', 'damaged_values': ['minor-damage', 'major-damage'],
        'destroyed_values': ['destroyed'], 'excluded_values': ['no-data (cloud)', 'no-data (no pre)'],
    },
    'impact': {
        'kind': 'points',
        'path': os.path.join(DONNEES, r"VENEZUELA\IMPACT Initiatives S1 DPM\impact_ven_earthquake_sentinel1_damaged_20260625_v3.geojson"),
        'class_field': None, 'damaged_values': None, 'destroyed_values': [],
    },
    'microsoft': {
        'kind': 'points',
        'path': os.path.join(DONNEES, r"VENEZUELA\Microsoft AI for Good Lab\from_huggingFace\venezuela_microsoft_damage_points.geojson"),
        'class_field': None, 'damaged_values': None, 'destroyed_values': [],
    },
    'osu': {
        'kind': 'points',
        'path': os.path.join(DONNEES, r"VENEZUELA\OSU S1 InSAR CCD\EMSR884_all_buildings_damage_assessment_20260701_v1.gpkg"),
        'class_field': 'label', 'damaged_values': ['likely_damaged'], 'destroyed_values': [],
        'excluded_values': ['not_assessed'],
    },
    'uhsail': {
        'kind': 'points',
        'path': os.path.join(DONNEES, r"VENEZUELA\UH SAIL\Building damaged.gpkg"),
        'class_field': 'grade', 'damaged_values': ['damaged'], 'destroyed_values': ['destroyed'],
    },
    'unep': {
        'kind': 'points',
        'path': os.path.join(DONNEES, r"VENEZUELA\UNEP\debris_buildings.gpkg"),
        'class_field': None, 'damaged_values': None, 'destroyed_values': [],
    },
}

RASTER_METHODS_ON_OVERTURE = {
    # (chemin, champ_classe, valeurs_endommage, valeurs_detruit) — PAR SOURCE,
    # chacune peut avoir un schema different (confirme par le pre-flight check).
    'eos':     (os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\raster_sources_on_overture\eos_on_overture.gpkg"),
                'classe', ['damaged'], ['destroyed']),
    'nasa_s1': (os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\raster_sources_on_overture\nasa_s1_on_overture.gpkg"),
                'classe', ['damaged'], ['destroyed']),
    'nasa_s2': (os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\raster_sources_on_overture\nasa_s2_on_overture.gpkg"),
                'classe', ['damaged'], ['destroyed']),
    'ungsc':   (os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\raster_sources_on_overture\ungsc_on_overture.gpkg"),
                'classe', ['damaged'], ['destroyed']),
    # t_test et bdpm suivent un autre schema : champ 'dmg_label', valeurs
    # 'Affected' / 'Not affected'. Le partage est binaire, sans distinction
    # entre endommage et detruit.
    't_test':  (os.path.join(DONNEES, r"VENEZUELA\MOI\T_test\AOI_CEMS\T_Test_Buildling_damage_overture.gpkg"),
                'dmg_label', ['Affected'], []),
    'bdpm':    (os.path.join(DONNEES, r"VENEZUELA\MOI\T_test\AOI_CEMS\BDPM_building_damage.gpkg"),
                'dmg_label', ['Affected'], []),
}

SOURCE_AOI_VECTOR_PATHS = {
    'chatmap':   os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\chatmap_aoi.gpkg"),
    'disha':     os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\disha_aoi.gpkg"),
    'ems':       os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\ems_aoi.gpkg"),
    'fair':      os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\fair_aoi.gpkg"),
    'impact':    os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\impact_aoi.gpkg"),
    'microsoft': os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\microsoft_aoi.gpkg"),
    'osu':       os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\osu_aoi.gpkg"),
    'uhsail':    os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\uhsail_aoi.gpkg"),
    'unep':      os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\unep_aoi.gpkg"),
    'eos':       os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\eos_aoi.gpkg"),
    'nasa_s1':   os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\nasa_s1_aoi.gpkg"),
    'nasa_s2':   os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\nasa_s2_aoi.gpkg"),
    'ungsc':     os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\ungsc_aoi.gpkg"),
}
SOURCE_AOI_RASTER_PATHS = {
    't_test': os.path.join(DONNEES, r"VENEZUELA\MOI\T_test\ALL_AOI_HOTOSM\T_stat_all_zones.tif"),
    'bdpm': os.path.join(DONNEES, r"VENEZUELA\MOI\BDPM\ALL_AOI_HOTOSM\coh_diff_fused.tif"),
}

OUT_DIR = os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\analyse")
OUT_CSV_MMI = os.path.join(OUT_DIR, "admin3_mmi_binsv3.csv")
OUT_CSV_DOMMAGES = os.path.join(OUT_DIR, "admin3_dommages_nuagesv3.csv")
OUT_XLSX = os.path.join(OUT_DIR, "synthese_admin3v3.xlsx")

SEP = "=" * 70

# =============================================================================
# HELPERS COMMUNS
# =============================================================================

WGS84 = osr.SpatialReference()
WGS84.ImportFromEPSG(4326)
WGS84.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)


def get_transform_to_wgs84(lyr):
    src_srs = lyr.GetSpatialRef()
    if src_srs is None:
        return None
    src_srs.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    if src_srs.IsSame(WGS84):
        return None
    return osr.CoordinateTransformation(src_srs, WGS84)


def load_union_geom(path):
    ds = ogr.Open(path)
    if ds is None:
        sys.exit(f"Impossible d'ouvrir : {path}")
    lyr = ds.GetLayer()
    tr = get_transform_to_wgs84(lyr)
    geoms = []
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        if not g2.IsValid():
            g2 = g2.MakeValid()
        geoms.append(g2)
    ds = None
    if not geoms:
        return None
    union = geoms[0]
    for g in geoms[1:]:
        union = union.Union(g)
    return union


def load_multi_union(paths):
    all_geoms = []
    for p in paths:
        if not os.path.exists(p):
            print(f"  Attention: fichier introuvable — {p}")
            continue
        g = load_union_geom(p)
        if g is not None:
            all_geoms.append(g)
    if not all_geoms:
        return None
    union = all_geoms[0]
    for g in all_geoms[1:]:
        union = union.Union(g)
    return union


def load_aoi_features_with_area(path, area_field):
    if area_field is None:
        return []
    ds = ogr.Open(path)
    if ds is None:
        sys.exit(f"Impossible d'ouvrir l'AOI : {path}")
    lyr = ds.GetLayer()
    defn = lyr.GetLayerDefn()
    field_names = [defn.GetFieldDefn(i).GetName() for i in range(defn.GetFieldCount())]
    if area_field not in field_names:
        sys.exit(f"Erreur: champ '{area_field}' introuvable dans {path}.\n"
                  f"Champs disponibles: {field_names}")
    tr = get_transform_to_wgs84(lyr)
    out = []
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        if not g2.IsValid():
            g2 = g2.MakeValid()
        out.append((g2, feat.GetField(area_field)))
    ds = None
    return out


def aoi_zones_for_geom(aoi_features, geom):
    labels = []
    for g, label in aoi_features:
        if g.Intersects(geom) and label not in labels:
            labels.append(label)
    return sorted(labels, key=lambda x: str(x)) if labels else [None]


def load_admin3_in_aoi(admin3_path, aoi_geom):
    """DECOUPE chaque admin3 exactement sur la limite de l'AOI (Intersection)
    — CHANGEMENT vs la version precedente qui gardait l'admin3 ENTIER. Tous
    les calculs (aire, batiments, population, MMI) portent desormais sur
    la geometrie DECOUPEE, pas sur l'admin3 complet."""
    ds = ogr.Open(admin3_path)
    if ds is None:
        sys.exit(f"Impossible d'ouvrir : {admin3_path}")
    lyr = ds.GetLayer()
    tr = get_transform_to_wgs84(lyr)
    field_names = [lyr.GetLayerDefn().GetFieldDefn(i).GetName()
                   for i in range(lyr.GetLayerDefn().GetFieldCount())]
    env = aoi_geom.GetEnvelope()
    lyr.SetSpatialFilterRect(env[0], env[2], env[1], env[3])
    out = []
    n_skipped = 0
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        if not g2.Intersects(aoi_geom):
            continue
        try:
            clipped = g2.Intersection(aoi_geom)
        except Exception:
            n_skipped += 1
            continue
        if clipped is None or clipped.IsEmpty():
            n_skipped += 1
            continue
        attrs = {fn: feat.GetField(fn) for fn in field_names}
        out.append({'geom': clipped, 'attrs': attrs})
    ds = None
    print(f"  {len(out)} admin3 dans l'AOI (DECOUPES sur la limite exacte de l'AOI)"
          + (f" — {n_skipped} ignore(s) (intersection vide/echec)" if n_skipped else ""))
    return out


def area_km2(geom):
    env = geom.GetEnvelope()
    lat_mid = (env[2] + env[3]) / 2
    deg_to_km_lon = 111.320 * np.cos(np.radians(lat_mid))
    deg_to_km_lat = 111.320
    return geom.GetArea() * deg_to_km_lon * deg_to_km_lat


def safe_intersection(geom_a, geom_b):
    try:
        if not geom_a.Intersects(geom_b):
            return None
        inter = geom_a.Intersection(geom_b)
        if inter is None or inter.IsEmpty():
            return None
        return inter
    except Exception:
        return None


def safe_difference(geom_a, geom_b):
    try:
        diff = geom_a.Difference(geom_b)
        if diff is None or diff.IsEmpty():
            return None
        return diff
    except Exception:
        return geom_a


def count_buildings_in_geom(buildings_path, geom, field_override=None):
    if geom is None:
        return 0
    ds = ogr.Open(buildings_path)
    if ds is None:
        return None
    lyr = ds.GetLayer()
    tr = get_transform_to_wgs84(lyr)
    env = geom.GetEnvelope()
    lyr.SetSpatialFilterRect(env[0], env[2], env[1], env[3])
    n = 0
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        c = g2.Centroid()
        if geom.Contains(c):
            n += 1
    ds = None
    return n


def raster_sum_in_geom(raster_path, geom):
    if geom is None:
        return 0.0
    ds = gdal.Open(raster_path)
    if ds is None:
        return None
    gt = ds.GetGeoTransform()
    prj = ds.GetProjection()
    nx_full, ny_full = ds.RasterXSize, ds.RasterYSize
    band = ds.GetRasterBand(1)
    nd = band.GetNoDataValue()

    src_srs = osr.SpatialReference(); src_srs.ImportFromWkt(prj)
    src_srs.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    tr_to_raster = None
    if not src_srs.IsSame(WGS84):
        tr_to_raster = osr.CoordinateTransformation(WGS84, src_srs)
    geom_r = geom.Clone()
    if tr_to_raster is not None:
        geom_r.Transform(tr_to_raster)

    env = geom_r.GetEnvelope()
    px0 = max(0, int((env[0] - gt[0]) / gt[1]) - 1)
    px1 = min(nx_full - 1, int((env[1] - gt[0]) / gt[1]) + 1)
    py0 = max(0, int((env[3] - gt[3]) / gt[5]) - 1)
    py1 = min(ny_full - 1, int((env[2] - gt[3]) / gt[5]) + 1)
    if px0 > px1 or py0 > py1:
        ds = None
        return 0.0

    win_nx, win_ny = px1 - px0 + 1, py1 - py0 + 1
    arr = band.ReadAsArray(px0, py0, win_nx, win_ny).astype(np.float64)
    if nd is not None:
        arr[np.isclose(arr, nd)] = np.nan
    ds = None

    win_gt = (gt[0] + px0 * gt[1], gt[1], 0, gt[3] + py0 * gt[5], 0, gt[5])
    mem_drv = gdal.GetDriverByName('MEM')
    mask_ds = mem_drv.Create('', win_nx, win_ny, 1, gdal.GDT_Byte)
    mask_ds.SetGeoTransform(win_gt)
    mask_ds.SetProjection(prj)

    mem_vec_drv = ogr.GetDriverByName('Memory')
    mem_vec_ds = mem_vec_drv.CreateDataSource('mask')
    mem_lyr = mem_vec_ds.CreateLayer('mask', srs=src_srs, geom_type=ogr.wkbMultiPolygon)
    feat = ogr.Feature(mem_lyr.GetLayerDefn())
    feat.SetGeometry(geom_r)
    mem_lyr.CreateFeature(feat)
    gdal.RasterizeLayer(mask_ds, [1], mem_lyr, burn_values=[1])
    mask_arr = mask_ds.GetRasterBand(1).ReadAsArray().astype(bool)
    mask_ds = None
    mem_vec_ds = None

    valid = mask_arr & np.isfinite(arr)
    return float(np.nansum(arr[valid])) if valid.any() else 0.0


def mmi_category(value):
    for lo, hi, label in MMI_BINS:
        if lo <= value < hi:
            return label
    return None


def load_mmi_cells_for_admin3(mmi_path, value_field, admin3_geom):
    ds = ogr.Open(mmi_path)
    if ds is None:
        sys.exit(f"Impossible d'ouvrir : {mmi_path}")
    lyr = ds.GetLayer()
    defn = lyr.GetLayerDefn()
    field_names = [defn.GetFieldDefn(i).GetName() for i in range(defn.GetFieldCount())]
    if value_field not in field_names:
        sys.exit(f"Erreur: champ '{value_field}' introuvable dans {mmi_path}.\n"
                  f"Champs disponibles: {field_names}")
    tr = get_transform_to_wgs84(lyr)
    env = admin3_geom.GetEnvelope()
    lyr.SetSpatialFilterRect(env[0], env[2], env[1], env[3])

    by_category = {}
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        if not g2.Intersects(admin3_geom):
            continue
        try:
            inter = g2.Intersection(admin3_geom)
        except Exception:
            continue
        if inter is None or inter.IsEmpty():
            continue
        val = feat.GetField(value_field)
        if val is None:
            continue
        cat = mmi_category(float(val))
        if cat is None:
            continue
        by_category.setdefault(cat, []).append(inter)
    ds = None

    out = {}
    for cat, geoms in by_category.items():
        multi = ogr.Geometry(ogr.wkbMultiPolygon)
        for g in geoms:
            if g.GetGeometryType() in (ogr.wkbMultiPolygon, ogr.wkbMultiPolygon25D):
                for i in range(g.GetGeometryCount()):
                    multi.AddGeometry(g.GetGeometryRef(i))
            elif g.GetGeometryType() in (ogr.wkbPolygon, ogr.wkbPolygon25D):
                multi.AddGeometry(g)
        out[cat] = multi
    return out


def count_damage_points(cfg, geom):
    path = cfg.get('path')
    if not path or path == '****' or not os.path.exists(path) or geom is None:
        return None, None
    ds = ogr.Open(path)
    if ds is None:
        return None, None
    lyr = ds.GetLayer()
    tr = get_transform_to_wgs84(lyr)
    env = geom.GetEnvelope()
    lyr.SetSpatialFilterRect(env[0], env[2], env[1], env[3])
    class_field = cfg.get('class_field')
    damaged_vals = cfg.get('damaged_values')
    destroyed_vals = cfg.get('destroyed_values') or []
    excluded_vals = cfg.get('excluded_values') or []
    n_damaged, n_destroyed = 0, 0
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        c = g2.Centroid()
        if not geom.Contains(c):
            continue
        if class_field is None:
            n_damaged += 1
            continue
        raw = feat.GetField(class_field)
        if raw in excluded_vals:
            continue
        if raw in destroyed_vals:
            n_destroyed += 1
        elif damaged_vals is None or raw in damaged_vals:
            n_damaged += 1
    ds = None
    return n_damaged, n_destroyed


def count_damage_csv(cfg, geom):
    path = cfg.get('path')
    if not path or path == '****' or not os.path.exists(path) or geom is None:
        return None, None
    class_field = cfg.get('class_field')
    damaged_vals = cfg.get('damaged_values')
    env = geom.GetEnvelope()
    n_damaged, n_destroyed = 0, 0
    with open(path, newline='', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            try:
                lon = float(row.get('lon') or row.get('longitude') or row.get('x'))
                lat = float(row.get('lat') or row.get('latitude') or row.get('y'))
            except (TypeError, ValueError):
                continue
            if not (env[0] <= lon <= env[1] and env[2] <= lat <= env[3]):
                continue
            pt = ogr.Geometry(ogr.wkbPoint)
            pt.AddPoint(lon, lat)
            if not geom.Contains(pt):
                continue
            raw = row.get(class_field)
            if damaged_vals is not None and raw in [str(v) for v in damaged_vals]:
                n_damaged += 1
    return n_damaged, n_destroyed


def count_damage_for_source(cfg, geom):
    kind = cfg.get('kind')
    if kind == 'points':
        return count_damage_points(cfg, geom)
    elif kind == 'points_csv':
        return count_damage_csv(cfg, geom)
    return None, None


def count_raster_method_on_overture(path, class_field, damaged_values, destroyed_values, geom):
    if geom is None or not os.path.exists(path):
        return None, None
    ds = ogr.Open(path)
    if ds is None:
        return None, None
    lyr = ds.GetLayer()
    defn = lyr.GetLayerDefn()
    field_names = [defn.GetFieldDefn(i).GetName() for i in range(defn.GetFieldCount())]
    if class_field not in field_names:
        sys.exit(f"Erreur: champ '{class_field}' introuvable dans {path}.\n"
                  f"Champs disponibles: {field_names}\n"
                  f"Corrige RASTER_METHODS_ON_OVERTURE en haut du script.")
    tr = get_transform_to_wgs84(lyr)
    env = geom.GetEnvelope()
    lyr.SetSpatialFilterRect(env[0], env[2], env[1], env[3])
    n_damaged, n_destroyed = 0, 0
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        c = g2.Centroid()
        if not geom.Contains(c):
            continue
        raw = feat.GetField(class_field)
        if raw in destroyed_values:
            n_destroyed += 1
        elif raw in damaged_values:
            n_damaged += 1
    ds = None
    return n_damaged, n_destroyed


def raster_bbox_polygon(path):
    ds = gdal.Open(path)
    if ds is None:
        return None
    gt = ds.GetGeoTransform()
    prj = ds.GetProjection()
    nx, ny = ds.RasterXSize, ds.RasterYSize
    ds = None
    corners = [(gt[0], gt[3]), (gt[0] + nx * gt[1], gt[3]),
               (gt[0] + nx * gt[1], gt[3] + ny * gt[5]), (gt[0], gt[3] + ny * gt[5])]
    srs = osr.SpatialReference(); srs.ImportFromWkt(prj)
    tr = None
    if not srs.IsSame(WGS84):
        srs.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
        tr = osr.CoordinateTransformation(srs, WGS84)
    ring = ogr.Geometry(ogr.wkbLinearRing)
    for x, y in corners:
        if tr is not None:
            x, y, _ = tr.TransformPoint(x, y)
        ring.AddPoint(x, y)
    ring.CloseRings()
    poly = ogr.Geometry(ogr.wkbPolygon)
    poly.AddGeometry(ring)
    return poly


def load_source_aoi(source_name):
    if source_name in SOURCE_AOI_VECTOR_PATHS:
        path = SOURCE_AOI_VECTOR_PATHS[source_name]
        if not os.path.exists(path):
            print(f"  Attention: AOI introuvable pour '{source_name}' — {path}")
            return None
        return load_union_geom(path)
    elif source_name in SOURCE_AOI_RASTER_PATHS:
        path = SOURCE_AOI_RASTER_PATHS[source_name]
        if not os.path.exists(path):
            print(f"  Attention: raster introuvable pour '{source_name}' — {path}")
            return None
        return raster_bbox_polygon(path)
    return None


# =============================================================================
# MAIN
# =============================================================================

def preflight_check_raster_methods():
    """Verifie TOUS les fichiers RASTER_METHODS_ON_OVERTURE (champ +
    valeurs, chacun avec SON PROPRE schema) AVANT de lancer la boucle
    admin3 couteuse — pour ne pas perdre le calcul MMI/dommages deja
    fait sur les premiers admin3 si un fichier a un schema inattendu."""
    print(f"{SEP}\n  Verification prealable des fichiers raster/Overture\n{SEP}")
    ok = True
    for name, (path, class_field, damaged_values, destroyed_values) in RASTER_METHODS_ON_OVERTURE.items():
        if not os.path.exists(path):
            print(f"  {name:10s} : ATTENTION fichier introuvable — {path}")
            continue
        ds = ogr.Open(path)
        if ds is None:
            print(f"  {name:10s} : ATTENTION impossible d'ouvrir le fichier")
            continue
        lyr = ds.GetLayer()
        defn = lyr.GetLayerDefn()
        field_names = [defn.GetFieldDefn(i).GetName() for i in range(defn.GetFieldCount())]
        if class_field not in field_names:
            print(f"  {name:10s} : ERREUR — champ '{class_field}' introuvable.")
            print(f"               Champs disponibles: {field_names}")
            ok = False
            ds = None
            continue
        vals = {}
        n_sample = 0
        n_total = lyr.GetFeatureCount()
        # Echantillonnage par PAS REGULIER sur tout le fichier — pas
        # juste les premieres lignes (deja pris en defaut : sur BDPM, les
        # 20000 premieres lignes ne contenaient presque aucun "Affected",
        # alors que le fichier complet en a 27% — donnee non uniformement
        # repartie, probablement triee par zone geographique/ID).
        stride = max(1, n_total // 20000) if n_total > 0 else 1
        for i, feat in enumerate(lyr):
            if i % stride != 0:
                continue
            n_sample += 1
            v = feat.GetField(class_field)
            vals[v] = vals.get(v, 0) + 1
        print(f"  {name:10s} : champ '{class_field}' — valeurs (echantillon {n_sample}): {vals}")
        matched_damaged = [v for v in vals if v in damaged_values]
        matched_destroyed = [v for v in vals if v in destroyed_values]
        if not matched_damaged and not matched_destroyed:
            print(f"               ATTENTION — AUCUNE des valeurs ci-dessus ne correspond a "
                  f"damaged_values={damaged_values} ni destroyed_values={destroyed_values} — "
                  f"le comptage donnera 0 partout. Corrige RASTER_METHODS_ON_OVERTURE['{name}'] avant de lancer.")
            ok = False
        ds = None
    if not ok:
        sys.exit("\nCorrige la configuration ci-dessus avant de relancer — "
                  "pas de calcul lance pour eviter de perdre du temps.")
    print("  OK — tous les fichiers raster/Overture sont valides.\n")


def main():
    os.makedirs(OUT_DIR, exist_ok=True)

    preflight_check_raster_methods()

    print(SEP)
    print(f"  Chargement de l'AOI (PARTAGEE) : {AOI_PATH}")
    print(SEP)
    aoi_geom = load_union_geom(AOI_PATH)
    if aoi_geom is None:
        sys.exit(f"AOI vide ou illisible : {AOI_PATH}")
    aoi_features = load_aoi_features_with_area(AOI_PATH, AOI_AREA_FIELD)

    print(f"\n{SEP}")
    print("  Admin3 dans l'AOI (PARTAGE — meme selection pour les deux analyses)")
    print(SEP)
    admin3_list = load_admin3_in_aoi(ADMIN3_PATH, aoi_geom)
    if not admin3_list:
        sys.exit("Aucun admin3 trouve dans cette AOI.")

    print(f"\n{SEP}")
    print("  Chargement des zones NON ANALYSEES (couverture nuageuse)")
    print(SEP)
    cloud_geom = load_multi_union(NOT_ANALYSED_PATHS)

    source_names = list(DAMAGE_CONFIG.keys()) + list(RASTER_METHODS_ON_OVERTURE.keys())
    print(f"\n{SEP}")
    print(f"  Prechargement des AOI par source ({len(source_names)} sources)")
    print(SEP)
    source_aois = {}
    for sn in source_names:
        g = load_source_aoi(sn)
        source_aois[sn] = g
        print(f"  {sn:12s} : {'OK' if g is not None else 'ECHEC'}")

    mmi_rows = []
    dommages_rows = []

    for i, unit in enumerate(admin3_list):
        g = unit['geom']
        admin3_id = unit['attrs'].get(ADMIN3_ID_FIELD, i)
        zones = aoi_zones_for_geom(aoi_features, g)
        admin3_area_km2 = area_km2(g)
        # 'g' est DEJA decoupe sur l'AOI (voir load_admin3_in_aoi) — plus
        # besoin de reintersecter avec aoi_geom, g y est deja entierement contenu.
        within_ems = g
        urban_area_km2 = (raster_sum_in_geom(GHS_BUILT_PATH, g) or 0.0) / 1e6

        if cloud_geom is not None:
            cloud_portion = safe_intersection(within_ems, cloud_geom)
            non_cloud_portion = safe_difference(within_ems, cloud_geom)
        else:
            cloud_portion = None
            non_cloud_portion = within_ems
        cloud_km2 = area_km2(cloud_portion) if cloud_portion is not None else 0.0
        cloud_pct_admin3 = 100 * cloud_km2 / admin3_area_km2 if admin3_area_km2 > 0 else None

        print(f"\n{SEP}")
        print(f"  [{i+1}/{len(admin3_list)}] {admin3_id} (zones={zones}) — "
              f"aire={admin3_area_km2:.2f}km2  urbain={urban_area_km2:.3f}km2")
        print(SEP)

        cells_by_cat = load_mmi_cells_for_admin3(MMI_SHP_PATH, MMI_VALUE_FIELD, g)
        mmi_results = []
        for _, _, cat in MMI_BINS:
            cat_geom = cells_by_cat.get(cat)
            if cat_geom is None or cat_geom.GetGeometryCount() == 0:
                area_val, n_ov, n_osm, pop = 0.0, 0, 0, 0.0
            else:
                area_val = area_km2(cat_geom)
                n_ov = count_buildings_in_geom(BUILDINGS_OVERTURE_PATH, cat_geom, BUILDINGS_OVERTURE_FIELD_OVERRIDE) or 0
                n_osm = count_buildings_in_geom(BUILDINGS_OSM_PATH, cat_geom, BUILDINGS_OSM_FIELD_OVERRIDE) or 0
                pop = raster_sum_in_geom(POPULATION_PATH, cat_geom) or 0.0
            mmi_results.append((cat, area_val, n_ov, n_osm, pop))
            print(f"    MMI {cat:8s} : aire={area_val:.2f}km2  overture={n_ov}  osm={n_osm}  pop={pop:.0f}")

        for zone in zones:
            for cat, area_val, n_ov, n_osm, pop in mmi_results:
                mmi_rows.append({
                    'admin3_id': admin3_id, 'aoi_zone': zone,
                    **unit['attrs'],
                    'mmi_categorie': cat, 'area_km2': round(area_val, 3),
                    'n_buildings_overture': n_ov, 'n_buildings_osm': n_osm,
                    'population': round(pop, 0),
                })

        total_buildings_by_zone_type = {
            'non_cloud': count_buildings_in_geom(BUILDINGS_OVERTURE_PATH, non_cloud_portion, BUILDINGS_OVERTURE_FIELD_OVERRIDE),
            'cloud': count_buildings_in_geom(BUILDINGS_OVERTURE_PATH, cloud_portion, BUILDINGS_OVERTURE_FIELD_OVERRIDE),
        }

        dommages_results = []
        for source_name in source_names:
            for zone_type, zone_geom in (('non_cloud', non_cloud_portion), ('cloud', cloud_portion)):
                if zone_geom is None:
                    n_damaged, n_destroyed = None, None
                elif source_name in DAMAGE_CONFIG:
                    n_damaged, n_destroyed = count_damage_for_source(DAMAGE_CONFIG[source_name], zone_geom)
                else:
                    rm_path, rm_class_field, rm_damaged, rm_destroyed = RASTER_METHODS_ON_OVERTURE[source_name]
                    n_damaged, n_destroyed = count_raster_method_on_overture(
                        rm_path, rm_class_field, rm_damaged, rm_destroyed, zone_geom)
                n_total = ((n_damaged or 0) + (n_destroyed or 0)
                           if (n_damaged is not None or n_destroyed is not None) else None)

                source_aoi_geom = source_aois.get(source_name)
                if source_aoi_geom is not None and zone_geom is not None:
                    inter = safe_intersection(source_aoi_geom, zone_geom)
                    source_aoi_area_km2 = area_km2(inter) if inter is not None else 0.0
                    n_analysed = count_buildings_in_geom(BUILDINGS_OVERTURE_PATH, inter, BUILDINGS_OVERTURE_FIELD_OVERRIDE)
                else:
                    source_aoi_area_km2 = None
                    n_analysed = 0

                total_b = total_buildings_by_zone_type.get(zone_type) or 0
                coverage = (n_analysed / total_b) if total_b > 0 and n_analysed is not None else None
                damage_fraction = (n_total / n_analysed) if (n_analysed and n_total is not None) else None

                dommages_results.append((source_name, zone_type, source_aoi_area_km2, total_b,
                                          n_analysed, coverage, n_damaged, n_destroyed, n_total, damage_fraction))

        for zone in zones:
            for (source_name, zone_type, source_aoi_area_km2, total_b, n_analysed, coverage,
                 n_damaged, n_destroyed, n_total, damage_fraction) in dommages_results:
                dommages_rows.append({
                    'admin3_id': admin3_id, 'aoi_zone': zone,
                    **unit['attrs'],
                    'admin3_area_km2': round(admin3_area_km2, 2),
                    'urban_area_km2': round(urban_area_km2, 3),
                    'cloud_area_km2': round(cloud_km2, 2),
                    'cloud_pct_admin3': round(cloud_pct_admin3, 1) if cloud_pct_admin3 is not None else None,
                    'source': source_name, 'zone_type': zone_type,
                    'source_aoi_area_km2': round(source_aoi_area_km2, 2) if source_aoi_area_km2 is not None else None,
                    'total_buildings': total_b, 'analysed': n_analysed,
                    'coverage': round(coverage, 4) if coverage is not None else None,
                    'n_endommages': n_damaged, 'n_detruits': n_destroyed, 'n_total': n_total,
                    'damage_fraction': round(damage_fraction, 4) if damage_fraction is not None else None,
                })

    with open(OUT_CSV_MMI, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(mmi_rows[0].keys()))
        w.writeheader(); w.writerows(mmi_rows)
    print(f"\nEcrit : {OUT_CSV_MMI}")

    with open(OUT_CSV_DOMMAGES, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(dommages_rows[0].keys()))
        w.writeheader(); w.writerows(dommages_rows)
    print(f"Ecrit : {OUT_CSV_DOMMAGES}")

    build_excel(mmi_rows, dommages_rows, source_names)
    print(f"Ecrit : {OUT_XLSX}")


# =============================================================================

def build_excel(mmi_rows, dommages_rows, source_names):
    wb = openpyxl.Workbook()
    FONT = "Arial"
    header_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="2E5C8A", end_color="2E5C8A", fill_type="solid")
    data_font = Font(name=FONT, size=10)
    title_font = Font(name=FONT, bold=True, size=13)

    def write_table(ws, rows, start_row=1):
        headers = list(rows[0].keys())
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=c, value=h)
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for r, row in enumerate(rows, start=start_row + 1):
            for c, h in enumerate(headers, start=1):
                cell = ws.cell(row=r, column=c, value=row.get(h))
                cell.font = data_font
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 16
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        return headers, start_row + 1, start_row + len(rows)

    ws_mmi = wb.active
    ws_mmi.title = "Donnees_MMI"
    mmi_headers, mmi_first, mmi_last = write_table(ws_mmi, mmi_rows)

    ws_dmg = wb.create_sheet("Donnees_Dommages")
    dmg_headers, dmg_first, dmg_last = write_table(ws_dmg, dommages_rows)

    def col_letter(headers, name):
        return get_column_letter(headers.index(name) + 1)

    ws_syn1 = wb.create_sheet("Synthese_Sources")
    ws_syn1["A1"] = "Synthese par source et par zone AOI"
    ws_syn1["A1"].font = title_font
    zones_dmg = sorted(set(r['aoi_zone'] for r in dommages_rows if r['aoi_zone']))
    headers1 = ["Source", "Zone", "total_buildings", "analysed", "n_endommages",
                "n_detruits", "n_total", "damage_fraction"]
    for c, h in enumerate(headers1, start=1):
        cell = ws_syn1.cell(row=3, column=c, value=h)
        cell.font = header_font; cell.fill = header_fill
    src_col = col_letter(dmg_headers, 'source')
    zone_col = col_letter(dmg_headers, 'aoi_zone')
    tb_col = col_letter(dmg_headers, 'total_buildings')
    an_col = col_letter(dmg_headers, 'analysed')
    nd_col = col_letter(dmg_headers, 'n_endommages')
    nde_col = col_letter(dmg_headers, 'n_detruits')
    nt_col = col_letter(dmg_headers, 'n_total')
    rng = f"Donnees_Dommages!${src_col}${dmg_first}:${src_col}${dmg_last}"
    rng_zone = f"Donnees_Dommages!${zone_col}${dmg_first}:${zone_col}${dmg_last}"
    r = 4
    for source_name in source_names:
        for zone in zones_dmg:
            ws_syn1.cell(row=r, column=1, value=source_name).font = data_font
            ws_syn1.cell(row=r, column=2, value=zone).font = data_font
            for col_idx, data_col in [(3, tb_col), (4, an_col), (5, nd_col), (6, nde_col), (7, nt_col)]:
                formula = (f'=SUMIFS(Donnees_Dommages!${data_col}${dmg_first}:${data_col}${dmg_last},'
                           f'{rng},$A{r},{rng_zone},$B{r})')
                ws_syn1.cell(row=r, column=col_idx, value=formula).font = data_font
            ws_syn1.cell(row=r, column=8, value=f'=IFERROR($G{r}/$D{r},0)').font = data_font
            r += 1
    for c in range(1, 9):
        ws_syn1.column_dimensions[get_column_letter(c)].width = 16

    ws_syn2 = wb.create_sheet("Synthese_MMI_Zone")
    ws_syn2["A1"] = "Synthese par zone AOI et par tranche MMI"
    ws_syn2["A1"].font = title_font
    zones_mmi = sorted(set(r['aoi_zone'] for r in mmi_rows if r['aoi_zone']))
    cats = [c for _, _, c in MMI_BINS]
    mzone_col = col_letter(mmi_headers, 'aoi_zone')
    mcat_col = col_letter(mmi_headers, 'mmi_categorie')
    marea_col = col_letter(mmi_headers, 'area_km2')
    mov_col = col_letter(mmi_headers, 'n_buildings_overture')
    mpop_col = col_letter(mmi_headers, 'population')
    rng_mzone = f"Donnees_MMI!${mzone_col}${mmi_first}:${mzone_col}${mmi_last}"
    rng_mcat = f"Donnees_MMI!${mcat_col}${mmi_first}:${mcat_col}${mmi_last}"

    row3 = ["AOI Zone"] + [f"Surface {c}" for c in cats] + [f"Batiments Overture {c}" for c in cats] + \
           [f"Population {c}" for c in cats]
    for c, h in enumerate(row3, start=1):
        cell = ws_syn2.cell(row=3, column=c, value=h)
        cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(wrap_text=True)
    r = 4
    for zone in zones_mmi:
        ws_syn2.cell(row=r, column=1, value=zone).font = data_font
        col = 2
        for data_col in (marea_col, mov_col, mpop_col):
            for cat in cats:
                formula = (f'=SUMIFS(Donnees_MMI!${data_col}${mmi_first}:${data_col}${mmi_last},'
                           f'{rng_mzone},$A{r},{rng_mcat},"{cat}")')
                ws_syn2.cell(row=r, column=col, value=formula).font = data_font
                col += 1
        r += 1
    for c in range(1, 2 + 3 * len(cats)):
        ws_syn2.column_dimensions[get_column_letter(c)].width = 14

    wb.save(OUT_XLSX)


if __name__ == '__main__':
    main()
