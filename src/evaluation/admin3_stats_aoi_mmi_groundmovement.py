# -*- coding: utf-8 -*-
"""Tableau des unités administratives de niveau 3 contenues dans une emprise donnée, avec bâti, surface urbaine, déplacement InSAR maximal et intensité MMI maximale.

Les chemins d'entree et de sortie sont declares dans le bloc CONFIGURATION
ci-dessous, relativement aux racines de donnees fixees dans `config.py`.
"""

import os
import sys
import csv
import numpy as np
from config import DONNEES

try:
    from osgeo import ogr, osr, gdal
except ImportError:
    sys.exit("osgeo (gdal/ogr) requis.")
try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl requis : pip install openpyxl --break-system-packages")

ogr.UseExceptions()
gdal.UseExceptions()

# =============================================================================
# CONFIGURATION
# =============================================================================

# Change ce chemin pour basculer entre l'AOI EOS et l'AOI EMS complete
#AOI_PATH = os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\original\eos_aoi.gpkg")
AOI_PATH = os.path.join(DONNEES, r"VENEZUELA\EMS\All_EMSR884.gpkg")

ADMIN3_PATH = os.path.join(DONNEES, r"VENEZUELA\HDX\ven_admin_boundaries.geojson\ven_admin3.geojson")
ADMIN3_ID_FIELD = 'adm3_pcode'
ADMIN3_PARENT_ADM2_FIELD = 'adm2_pcode'  # pour l'heritage MMI

GHS_BUILT_PATH = os.path.join(DONNEES, r"VENEZUELA\HDX\GHS_BUILT_S_E2025_GLOBE_R2023A_54009_100_V1_0_R8_C12\GHS_BUILT_S_E2025_GLOBE_R2023A_54009_100_V1_0_R8_C12.tif")
POPULATION_PATH = os.path.join(DONNEES, r"VENEZUELA\HDX\ven_pop_2026_CN_100m_R2025A_v1.tif")

BUILDINGS_OVERTURE_PATH = os.path.join(DONNEES, r"VENEZUELA\Building\hot_eq_overture_ven_buildings_overture_gpkg\buildings.gpkg")
BUILDINGS_OVERTURE_FIELD_OVERRIDE = [
    'id', 'name', 'name_en', 'class', 'subtype', 'height', 'level',
    'num_floors', 'source', 'adm0_pcode', 'adm0_name', 'adm1_pcode',
    'adm1_name', 'adm2_pcode', 'adm2_name', 'adm3_pcode', 'adm3_name',
    'adm4_pcode', 'adm4_name',
]
BUILDINGS_OSM_PATH = os.path.join(DONNEES, r"VENEZUELA\Building\hot_eq_osm_ven_buildings_osm_gpkg\buildings.gpkg")
BUILDINGS_OSM_FIELD_OVERRIDE = [
    'id', 'name', 'name_en', 'name_es', 'building', 'building_levels',
    'building_materials', 'addr_full', 'addr_housenumber', 'addr_street',
    'addr_city', 'office', 'source', 'adm0_pcode', 'adm0_name',
    'adm1_pcode', 'adm1_name', 'adm2_pcode', 'adm2_name', 'adm3_pcode',
    'adm3_name', 'adm4_pcode', 'adm4_name', 'name_latin',
]

# Mouvement de sol (deplacement InSAR LOS) — verifie sur le fichier reel
GROUND_MOVEMENT_PATH = os.path.join(DONNEES, r"VENEZUELA\EMS\EMSR884_products\EMSR884_ALL\EMSR884_groundMovement.gpkg")
GROUND_MOVEMENT_FIELD = "value"  # confirme : plages textuelles ("0.2 to 0.5", etc.)

# Intensite sismique MMI (UNOSAT) — au niveau admin2, jamais admin3
UNOSAT_XLSX_PATH = os.path.join(DONNEES, r"VENEZUELA\UNOSAT\UNOSAT Live webmap - M 7.5-Caracasearthquake(24June2026at22_05_11UTC).xlsx")
UNOSAT_SHEET = "Data"
UNOSAT_ADM2_FIELD = "adm2_pcode"
UNOSAT_ZONE_FIELD = "Zone"

# Ordre de gravite MMI — plus le chiffre est haut, plus c'est intense
MMI_ORDER = {
    'II - III (Weak)': 1,
    'IV (Light)': 2,
    'V (Moderate)': 3,
    'VI (Strong)': 4,
    'VII (Very Strong)': 5,
    'VIII (Severe)': 6,
    'IX (Violent)': 7,
}

OUT_CSV = os.path.join(DONNEES, r"VENEZUELA\MOI\Source_AOI\analyse\CEMSALL_admin3_dans_aoi.csv")

SEP = "=" * 70

# =============================================================================
# HELPERS
# =============================================================================

WGS84 = osr.SpatialReference()
WGS84.ImportFromEPSG(4326)
WGS84.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)


def get_transform_to_wgs84(lyr):
    src_srs = lyr.GetSpatialRef()
    if src_srs is None:
        return None
    src_srs.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    if src_srs.IsSame(WGS84):
        return None
    return osr.CoordinateTransformation(src_srs, WGS84)


def load_aoi_union(path):
    ds = ogr.Open(path)
    if ds is None:
        sys.exit(f"Impossible d'ouvrir l'AOI : {path}")
    lyr = ds.GetLayer()
    tr = get_transform_to_wgs84(lyr)
    geoms = []
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        if not g2.IsValid():
            g2 = g2.MakeValid()
        geoms.append(g2)
    ds = None
    if not geoms:
        sys.exit(f"AOI vide : {path}")
    if len(geoms) == 1:
        return geoms[0]
    multi = ogr.Geometry(ogr.wkbMultiPolygon)
    for g in geoms:
        if g.GetGeometryType() in (ogr.wkbMultiPolygon, ogr.wkbMultiPolygon25D):
            for i in range(g.GetGeometryCount()):
                multi.AddGeometry(g.GetGeometryRef(i))
        else:
            multi.AddGeometry(g)
    return multi


def load_admin3_in_aoi(admin3_path, aoi_geom):
    ds = ogr.Open(admin3_path)
    if ds is None:
        sys.exit(f"Impossible d'ouvrir : {admin3_path}")
    lyr = ds.GetLayer()
    tr = get_transform_to_wgs84(lyr)
    field_names = [lyr.GetLayerDefn().GetFieldDefn(i).GetName()
                   for i in range(lyr.GetLayerDefn().GetFieldCount())]

    env = aoi_geom.GetEnvelope()
    lyr.SetSpatialFilterRect(env[0], env[2], env[1], env[3])

    out = []
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        if not g2.Intersects(aoi_geom):
            continue
        attrs = {fn: feat.GetField(fn) for fn in field_names}
        out.append({'geom': g2, 'attrs': attrs})
    ds = None
    print(f"  {len(out)} admin3 dans l'AOI (entiers, non decoupes)")
    return out


def area_km2(geom):
    env = geom.GetEnvelope()
    lat_mid = (env[2] + env[3]) / 2
    deg_to_km_lon = 111.320 * np.cos(np.radians(lat_mid))
    deg_to_km_lat = 111.320
    return geom.GetArea() * deg_to_km_lon * deg_to_km_lat


def raster_sum_in_geom(raster_path, geom):
    ds = gdal.Open(raster_path)
    if ds is None:
        return None
    gt = ds.GetGeoTransform()
    prj = ds.GetProjection()
    nx_full, ny_full = ds.RasterXSize, ds.RasterYSize
    band = ds.GetRasterBand(1)
    nd = band.GetNoDataValue()

    src_srs = osr.SpatialReference(); src_srs.ImportFromWkt(prj)
    src_srs.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    tr_to_raster = None
    if not src_srs.IsSame(WGS84):
        tr_to_raster = osr.CoordinateTransformation(WGS84, src_srs)
    geom_r = geom.Clone()
    if tr_to_raster is not None:
        geom_r.Transform(tr_to_raster)

    env = geom_r.GetEnvelope()
    px0 = max(0, int((env[0] - gt[0]) / gt[1]) - 1)
    px1 = min(nx_full - 1, int((env[1] - gt[0]) / gt[1]) + 1)
    py0 = max(0, int((env[3] - gt[3]) / gt[5]) - 1)
    py1 = min(ny_full - 1, int((env[2] - gt[3]) / gt[5]) + 1)
    if px0 > px1 or py0 > py1:
        ds = None
        return 0.0

    win_nx, win_ny = px1 - px0 + 1, py1 - py0 + 1
    arr = band.ReadAsArray(px0, py0, win_nx, win_ny).astype(np.float64)
    if nd is not None:
        arr[np.isclose(arr, nd)] = np.nan
    ds = None

    win_gt = (gt[0] + px0 * gt[1], gt[1], 0, gt[3] + py0 * gt[5], 0, gt[5])
    mem_drv = gdal.GetDriverByName('MEM')
    mask_ds = mem_drv.Create('', win_nx, win_ny, 1, gdal.GDT_Byte)
    mask_ds.SetGeoTransform(win_gt)
    mask_ds.SetProjection(prj)

    mem_vec_drv = ogr.GetDriverByName('Memory')
    mem_vec_ds = mem_vec_drv.CreateDataSource('mask')
    mem_lyr = mem_vec_ds.CreateLayer('mask', srs=src_srs, geom_type=ogr.wkbMultiPolygon)
    feat = ogr.Feature(mem_lyr.GetLayerDefn())
    feat.SetGeometry(geom_r)
    mem_lyr.CreateFeature(feat)
    gdal.RasterizeLayer(mask_ds, [1], mem_lyr, burn_values=[1])
    mask_arr = mask_ds.GetRasterBand(1).ReadAsArray().astype(bool)
    mask_ds = None
    mem_vec_ds = None

    valid = mask_arr & np.isfinite(arr)
    return float(np.nansum(arr[valid])) if valid.any() else 0.0


def count_buildings_in_geom(buildings_path, geom, field_override):
    ds = ogr.Open(buildings_path)
    if ds is None:
        return None
    lyr = ds.GetLayer()
    tr = get_transform_to_wgs84(lyr)
    env = geom.GetEnvelope()
    lyr.SetSpatialFilterRect(env[0], env[2], env[1], env[3])
    n = 0
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        c = g2.Centroid()
        if geom.Contains(c):
            n += 1
    ds = None
    return n


def parse_ground_movement_value(s):
    """Convertit une plage textuelle ('0.2 to 0.5', '-0.1 to -0.05',
    'Above 0.5') en une valeur representative — la borne la PLUS ELOIGNEE
    de zero (le deplacement le plus extreme dans cette plage)."""
    s = s.strip()
    if s.lower().startswith('above'):
        try:
            floor = float(s.split()[1])
        except (IndexError, ValueError):
            return None
        return floor + 0.1  # approximation — pas de vraie borne haute connue
    parts = s.split(' to ')
    if len(parts) != 2:
        return None
    try:
        a, b = float(parts[0]), float(parts[1])
    except ValueError:
        return None
    return a if abs(a) > abs(b) else b


def load_ground_movement(path, field):
    ds = ogr.Open(path)
    if ds is None:
        sys.exit(f"Impossible d'ouvrir : {path}")
    lyr = ds.GetLayer()
    defn = lyr.GetLayerDefn()
    field_names = [defn.GetFieldDefn(i).GetName() for i in range(defn.GetFieldCount())]
    if field not in field_names:
        sys.exit(f"Erreur: champ '{field}' introuvable dans {path}.\nChamps disponibles: {field_names}")
    tr = get_transform_to_wgs84(lyr)
    out = []
    n_unparsed = 0
    for feat in lyr:
        g = feat.GetGeometryRef()
        if g is None:
            continue
        g2 = g.Clone()
        if tr is not None:
            g2.Transform(tr)
        raw = feat.GetField(field)
        val = parse_ground_movement_value(raw) if raw else None
        if val is None:
            n_unparsed += 1
            continue
        out.append((g2, val))
    ds = None
    print(f"  Mouvement de sol : {len(out)} polygone(s) charge(s), {n_unparsed} non-parses")
    return out


def ground_movement_max_in_geom(gm_data, geom):
    env = geom.GetEnvelope()
    max_abs, max_val = None, None
    for g, val in gm_data:
        genv = g.GetEnvelope()
        if genv[1] < env[0] or genv[0] > env[1] or genv[3] < env[2] or genv[2] > env[3]:
            continue  # rejet rapide par bbox avant le test couteux
        if not g.Intersects(geom):
            continue
        if max_abs is None or abs(val) > max_abs:
            max_abs = abs(val)
            max_val = val
    return max_val


def load_unosat_mmi_by_adm2(xlsx_path, sheet, adm2_field, zone_field):
    """Retourne {adm2_pcode: zone_mmi_la_plus_intense_presente}."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if adm2_field not in headers or zone_field not in headers:
        sys.exit(f"Erreur: champ '{adm2_field}' ou '{zone_field}' introuvable dans "
                  f"la feuille '{sheet}'.\nChamps disponibles: {headers}")
    adm2_idx = headers.index(adm2_field)
    zone_idx = headers.index(zone_field)

    best_by_adm2 = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        adm2 = row[adm2_idx]
        zone = row[zone_idx]
        if adm2 is None or zone is None:
            continue
        rank = MMI_ORDER.get(zone)
        if rank is None:
            continue
        if adm2 not in best_by_adm2 or rank > MMI_ORDER.get(best_by_adm2[adm2], -1):
            best_by_adm2[adm2] = zone
    print(f"  UNOSAT : {len(best_by_adm2)} admin2 avec une zone MMI identifiee")
    return best_by_adm2


# =============================================================================
# MAIN
# =============================================================================

def main():
    print(SEP)
    print(f"  Chargement de l'AOI : {AOI_PATH}")
    print(SEP)
    aoi_geom = load_aoi_union(AOI_PATH)

    print(f"\n{SEP}")
    print("  Admin3 dans cette AOI")
    print(SEP)
    admin3_list = load_admin3_in_aoi(ADMIN3_PATH, aoi_geom)
    if not admin3_list:
        sys.exit("Aucun admin3 trouve dans cette AOI.")

    print(f"\n{SEP}")
    print("  Chargement du mouvement de sol")
    print(SEP)
    gm_data = load_ground_movement(GROUND_MOVEMENT_PATH, GROUND_MOVEMENT_FIELD)

    print(f"\n{SEP}")
    print("  Chargement de l'intensite MMI (UNOSAT, niveau admin2)")
    print(SEP)
    mmi_by_adm2 = load_unosat_mmi_by_adm2(UNOSAT_XLSX_PATH, UNOSAT_SHEET, UNOSAT_ADM2_FIELD, UNOSAT_ZONE_FIELD)

    print(f"\n{SEP}")
    print("  Calcul des statistiques par admin3")
    print(SEP)
    rows = []
    for i, unit in enumerate(admin3_list):
        g = unit['geom']
        a_km2 = area_km2(g)
        urban_m2 = raster_sum_in_geom(GHS_BUILT_PATH, g)
        urban_km2 = (urban_m2 or 0.0) / 1e6
        population = raster_sum_in_geom(POPULATION_PATH, g)
        n_overture = count_buildings_in_geom(BUILDINGS_OVERTURE_PATH, g, BUILDINGS_OVERTURE_FIELD_OVERRIDE)
        n_osm = count_buildings_in_geom(BUILDINGS_OSM_PATH, g, BUILDINGS_OSM_FIELD_OVERRIDE)
        gm_max = ground_movement_max_in_geom(gm_data, g)

        adm2_pcode = unit['attrs'].get(ADMIN3_PARENT_ADM2_FIELD)
        mmi_zone = mmi_by_adm2.get(adm2_pcode)

        admin3_id = unit['attrs'].get(ADMIN3_ID_FIELD, i)
        row = {
            'admin3_id': admin3_id,
            **unit['attrs'],
            'area_km2': round(a_km2, 2),
            'urban_area_km2': round(urban_km2, 2),
            'population': round(population, 0) if population is not None else None,
            'n_buildings_overture': n_overture,
            'n_buildings_osm': n_osm,
            'ground_movement_max_m': round(gm_max, 3) if gm_max is not None else None,
            'mmi_intensity_max_via_adm2': mmi_zone,
        }
        rows.append(row)
        print(f"  [{i+1}/{len(admin3_list)}] {admin3_id} — urbain={urban_km2:.2f}km2  "
              f"pop={(population or 0):.0f}  overture={n_overture}  osm={n_osm}  mvt_max={gm_max}  mmi={mmi_zone}")

    os.makedirs(os.path.dirname(OUT_CSV), exist_ok=True)
    fieldnames = list(rows[0].keys())
    with open(OUT_CSV, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)
    print(f"\nEcrit : {OUT_CSV}")


if __name__ == '__main__':
    main()
